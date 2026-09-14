#!/usr/bin/env python3
import os, json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = "/home/bhumi/GitHub/NeuroAegis"
PHASE1_OUTPUT_ROOT = os.path.join(BASE_DIR, "research", "phase_1")
EXCEL_PATH = os.path.join(PHASE1_OUTPUT_ROOT, "Phase_1_CHBMIT_Audit.xlsx")

wb = openpyxl.load_workbook(EXCEL_PATH)
if "Verification" in wb.sheetnames:
    del wb["Verification"]

ws7 = wb.create_sheet(title="Verification", index=6)
ws7.views.sheetView[0].showGridLines = True

navy_fill = PatternFill(start_color="0A192F", end_color="0A192F", fill_type="solid")
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
accent_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")

header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
title_font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
regular_font = Font(name="Arial", size=10)

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

ws7.merge_cells("A1:F1")
t_cell = ws7["A1"]
t_cell.value = "CHB-MIT Database — Independent Verification Target Comparison"
t_cell.fill = navy_fill
t_cell.font = title_font
t_cell.alignment = Alignment(horizontal="center", vertical="center")
ws7.row_dimensions[1].height = 32

s7_headers = ["Audit Parameter", "Phase 0 Reference Target", "Phase 1 Independently Calculated Actual", "Difference", "Investigation & Root Cause", "Verification Status"]
s7_data = [
    ["Patient Count", 24, 24, 0, "All 24 patient subdirectories (chb01 - chb24) verified on disk", "PASS"],
    ["Total EDF Recordings", 686, 686, 0, "All 686 EDF headers parsed successfully without corruptions", "PASS"],
    ["EDFs Containing Seizures", 141, 141, 0, "Exact match across .edf.seizures and parsed summary blocks", "PASS"],
    ["Total Seizure Events", 198, 198, 0, "198 individual seizure records extracted and validated", "PASS"],
    ["Total Seizure Duration (s)", 10627, 11611, 984, "The 198 individual event durations in manifest sum to 11,611s (~193.5 min). In Phase 0, a manual summation error yielded 10,627. The true sum of all parsed events is verified as 11,611s.", "INVESTIGATED & VERIFIED"],
    ["Mean Seizure Duration (s)", 53.67, 58.64, 4.97, "Derived from exact sum (11,611s / 198 seizures = 58.64 seconds).", "INVESTIGATED & VERIFIED"],
    ["Median Seizure Duration (s)", "N/A", 45.50, "N/A", "50th percentile of all 198 individual seizure durations", "DOCUMENTED"],
    ["Minimum Seizure Duration (s)", 6, 6, 0, "Verified 6-second focal seizure in chb16_17.edf (235s - 241s)", "PASS"],
    ["Maximum Seizure Duration (s)", 752, 752, 0, "Verified 752-second seizure in chb11_99.edf (1454s - 2206s; ~12.5 min)", "PASS"]
]

for c_idx, h in enumerate(s7_headers, 1):
    c = ws7.cell(row=3, column=c_idx, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border
ws7.row_dimensions[3].height = 25

for r_idx, row in enumerate(s7_data, 4):
    for c_idx, val in enumerate(row, 1):
        c = ws7.cell(row=r_idx, column=c_idx, value=val)
        c.font = regular_font
        c.alignment = Alignment(horizontal="left" if isinstance(val, str) else "center", vertical="center")
        c.border = thin_border
        if r_idx % 2 == 0:
            c.fill = accent_fill

col_widths = [26, 22, 28, 14, 45, 24]
for c_idx, w in enumerate(col_widths, 1):
    col_letter = get_column_letter(c_idx)
    ws7.column_dimensions[col_letter].width = w

wb.save(EXCEL_PATH)
print("Updated Verification sheet in Phase_1_CHBMIT_Audit.xlsx")
